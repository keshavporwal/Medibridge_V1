"""Generate MediBridge product backlog (consolidated 22 stories) as a styled .xlsx workbook.

Scope model (per project brief):
  * Frontend = HTML/CSS/vanilla JS, standalone prototype (localStorage), presented on its own.
  * Backend  = vanilla Java + JDBC + PostgreSQL, TERMINAL/console app, presented on its own.
  * Frontend and backend are NOT wired together and share no live connection.

Sheets: Cover | User Stories | Subtasks (grouped) | Epics | Scope & Assumptions | Backend DB Schema
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- palette
NAVY = "1F3864"
NAVY_MED = "2E5496"
BAND = "F2F5FB"
EPIC_FILL_C = "D9E1F2"
GREEN = "D9EAD3"
GREY = "EFEFEF"
RED = "F8D7DA"
AMBER = "FCE5CD"
FE = "DCE6F5"
BE = "E6DAF0"
DB = "FCE8D5"

HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, color=NAVY, size=20)
BOLD = Font(bold=True)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(vertical="top", horizontal="center", wrap_text=True)
thin = Side(style="thin", color="C9D3E5")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

PRIO_FILL = {"Must": RED, "Should": AMBER, "Could": GREY}
STATUS_FILL = {"Done": GREEN, "To Do": GREY, "In Progress": AMBER}
LAYER_FILL = {"Frontend": FE, "Backend": BE, "DB": DB}


def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)


def style_header(ws, ncols, row=1, color=NAVY):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill(color)
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 26


# =========================================================================== data
# Each story: id, epic, layer, role, title, acceptance, priority, points, status,
#             wireframes (list), subtasks (list of (text, est_hours))
S = [
    dict(id="US-01", epic="E1 Authentication & Access", layer="Frontend", role="All",
         title="Authentication & login",
         acc="Role selection updates imagery; valid creds route to role home; invalid show inline error; admin via dedicated portal.",
         prio="Must", pts=8, status="Done", wf=["US-01-login-register.html"],
         subs=[("Split-screen layout + role selector (Doctor/Patient/Receptionist) with role-based imagery", 3),
               ("Email/password sign-in with inline validation", 2),
               ("Role-based redirect to correct home on success", 2),
               ("Admin sign-in portal link + visual differentiation + default admin creds", 3),
               ("Session persistence + requireAuth guard + logout", 3)]),
    dict(id="US-02", epic="E1 Authentication & Access", layer="Frontend", role="Doctor/Patient/Staff",
         title="Registration & security questions",
         acc="Step 1 collects name/email/password (+clinic code for staff); step 2 sets 3 hashed questions; admin cannot self-register.",
         prio="Must", pts=8, status="Done", wf=["US-02-register-security-questions.html", "US-01-login-register.html"],
         subs=[("Step 1 form: full name, email, password with required markers", 3),
               ("Clinic-code field + validation against enrolled clinics (doctor/staff)", 2),
               ("Password strength meter & policy enforcement", 2),
               ("Step 2: choose & answer 3 distinct security questions", 2),
               ("SHA-256 salted hashing of answers (no plain-text)", 2),
               ("Block admin self-registration", 1)]),
    dict(id="US-03", epic="E1 Authentication & Access", layer="Frontend", role="All",
         title="Account recovery",
         acc="Identify by email; verify 3 questions (case-insensitive); lock after 3 fails; reset password with strength meter.",
         prio="Must", pts=5, status="Done",
         wf=["US-03-forgot-password-verify.html", "US-03-forgot-password-reset.html"],
         subs=[("Identify account by registered email", 1),
               ("Verify 3 security questions (case-insensitive hash match)", 2),
               ("Lock flow after 3 failed attempts", 1),
               ("Set new password with strength meter + confirm match", 2),
               ("Sign-out-all-sessions toggle", 1)]),
    dict(id="US-04", epic="E2 Shared UI Platform", layer="Frontend", role="All",
         title="Shared UI platform & account",
         acc="Reusable store, modals, validation, toasts, breadcrumbs, design system, and the account & security page.",
         prio="Must", pts=8, status="Done", wf=["US-04-account.html"],
         subs=[("localStorage data Store module (versioned CRUD)", 3),
               ("Modal engine (overlay/Esc/close, scrollable responsive footers)", 2),
               ("Form validation helpers (required *, phone, name, length, pattern, live-clear)", 3),
               ("Toasts + insight popovers", 2),
               ("Breadcrumb navigation on all pages", 2),
               ("Design-system tokens & components; responsive desktop-first", 3),
               ("Account & security page (profile info + change password)", 2)]),
    dict(id="US-05", epic="E3 Institute Admin", layer="Frontend", role="Admin",
         title="Admin dashboard & staff governance",
         acc="KPIs, volume trend, clinics quick view, staff table; onboard and deboard staff via modals.",
         prio="Must", pts=8, status="Done", wf=["US-05-admin-dashboard.html"],
         subs=[("Command-center KPIs (clinics/facilities/staff/volume) + insight popovers", 3),
               ("Patient volume trends chart", 2),
               ("Enrolled clinics quick directory", 2),
               ("Staff governance table with status (Active/On Leave/Inactive)", 2),
               ("Onboard Staff modal (name, email, clinic, role)", 3),
               ("Deboard Staff confirm modal (reassign consultations + revoke)", 2)]),
    dict(id="US-06", epic="E3 Institute Admin", layer="Frontend", role="Admin",
         title="Clinic enrollment & directory",
         acc="Directory lists all clinics; enroll modal captures name/specialty/address and generates a unique clinic code.",
         prio="Must", pts=5, status="Done", wf=["US-06-admin-clinics.html"],
         subs=[("Clinics directory page (specialty, code, staff count, status)", 2),
               ("Enroll New Clinic modal (name, specialty, address) + validation", 3),
               ("Generate & display unique clinic code", 2)]),
    dict(id="US-07", epic="E4 Doctor Workspace", layer="Frontend", role="Doctor",
         title="Doctor dashboard & document requests",
         acc="Bento dashboard with roster, daily schedule, request tracker; privacy-preserving document request modal.",
         prio="Must", pts=8, status="Done", wf=["US-07-doctor-dashboard.html"],
         subs=[("Bento dashboard KPIs", 2),
               ("Today's patient roster with Open Record", 2),
               ("Daily schedule gantt timeline", 2),
               ("Request Clinical Document modal (no vault browsing)", 3),
               ("Document request tracker + Shared-with-you list + preview", 3),
               ("Set weekly availability modal", 2)]),
    dict(id="US-08", epic="E4 Doctor Workspace", layer="Frontend", role="Doctor",
         title="Doctor scheduling (weekly & monthly)",
         acc="Weekly gantt shows slots/breaks/appointments; monthly grid with switcher; availability configurable.",
         prio="Must", pts=8, status="Done",
         wf=["US-08-doctor-schedule-weekly.html", "US-08-doctor-schedule-monthly.html"],
         subs=[("Weekly gantt: slots, breaks, appointment bars (visible w/o availability)", 3),
               ("Next Up queue", 2),
               ("Monthly calendar grid with month switcher", 3),
               ("Set availability modal (working days, shift, lunch)", 2)]),
    dict(id="US-09", epic="E4 Doctor Workspace", layer="Frontend", role="Doctor",
         title="Doctor patient records & notes",
         acc="Searchable directory; scoped clinical chart; consented shared docs only; log consultation notes; request docs.",
         prio="Must", pts=8, status="Done",
         wf=["US-09-doctor-patients.html", "US-09-doctor-patient-chart.html"],
         subs=[("Patient directory with search", 2),
               ("Clinical chart: summary, history, allergies, upcoming (scoped)", 3),
               ("Show only consented shared documents", 2),
               ("Log consultation note", 2),
               ("Request document modal from chart", 2)]),
    dict(id="US-10", epic="E5 Receptionist", layer="Frontend", role="Receptionist",
         title="Receptionist intake, triage & directory",
         acc="Intake KPIs, pending queue, availability roster; assign slots via review modal; patient directory.",
         prio="Must", pts=8, status="Done",
         wf=["US-10-receptionist-dashboard.html", "US-10-receptionist-patients.html"],
         subs=[("Intake KPIs (scheduled/pending/on-duty)", 2),
               ("Pending appointment requests queue", 2),
               ("Doctor availability roster", 2),
               ("Today's schedule list", 1),
               ("Review & Schedule Request modal (assign doctor + slot, dispatch)", 3),
               ("Doctor schedule inspection modal", 2),
               ("Patient directory", 2)]),
    dict(id="US-11", epic="E6 Patient Experience", layer="Frontend", role="Patient",
         title="Patient Health Hub",
         acc="Overview, consent panel, timeline; request appointment (doctor-exists guard) and grant document access.",
         prio="Must", pts=8, status="Done", wf=["US-11-patient-hub.html"],
         subs=[("Overview + mini-stats + upcoming visit card", 2),
               ("Consent & permissions panel (revoke)", 2),
               ("Vault mini + my appointment requests list", 2),
               ("Care & consultation timeline (gantt)", 3),
               ("Request Appointment modal (doctor-exists guard)", 3),
               ("Grant Document Access modal (files + duration)", 3)]),
    dict(id="US-12", epic="E6 Patient Experience", layer="Frontend", role="Patient",
         title="Patient appointments",
         acc="Calendar with appointment dots, upcoming and pending lists, and a request-appointment modal.",
         prio="Should", pts=5, status="Done", wf=["US-12-patient-appointments.html"],
         subs=[("Month calendar with appointment dots + switcher", 2),
               ("Upcoming & pending requests lists", 2),
               ("Request Appointment modal", 2)]),
    dict(id="US-13", epic="E6 Patient Experience", layer="Frontend", role="Patient",
         title="Patient document vault & consent",
         acc="Real-file upload (4MB cap), document list, grant-access modal, and PDF/image preview.",
         prio="Must", pts=8, status="Done", wf=["US-13-patient-vault.html"],
         subs=[("Drag & drop upload with real file storage (4MB cap)", 3),
               ("My documents list (type/date, preview/share)", 2),
               ("Pending requests + active access panels", 2),
               ("Grant Document Access modal", 2),
               ("Document preview (PDF/image, download fallback)", 3)]),
    dict(id="US-14", epic="E6 Patient Experience", layer="Frontend", role="Patient/Public",
         title="Patient digital ID & QR sharing",
         acc="ID pass with contact/medical info (validated phone); from-scratch QR encodes a public read-only share URL.",
         prio="Should", pts=8, status="Done",
         wf=["US-14-patient-profile.html", "US-14-patient-share.html"],
         subs=[("Digital ID pass page + contact/medical info (phone validation)", 2),
               ("From-scratch QR code generator (SVG)", 3),
               ("QR encodes shareable URL", 2),
               ("Public read-only share pass page (no login)", 2)]),
    dict(id="US-15", epic="E7 Public & Docs", layer="Frontend", role="Public/All",
         title="Public landing, PWA & wireframes",
         acc="Landing page communicates value + roles; PWA assets present; low-fi wireframes for all pages & popups.",
         prio="Should", pts=5, status="Done", wf=["US-15-landing.html", "index.html"],
         subs=[("Landing page (hero, roles, privacy band, footer)", 3),
               ("PWA manifest & app icons", 2),
               ("Wireframe documentation set (all pages + popups)", 3)]),
    # ---- Backend ----
    dict(id="US-16", epic="E8 Backend Database", layer="DB", role="Backend",
         title="Database foundation (Postgres + JDBC)",
         acc="Normalized schema created by idempotent DDL; JDBC connectivity works; default admin seeded.",
         prio="Must", pts=13, status="To Do", wf=[],
         subs=[("Design normalized Postgres schema (all entities)", 5),
               ("Write idempotent DDL (tables, enums, indexes)", 5),
               ("JDBC connection utility (config-driven)", 3),
               ("Seed default admin row (hashed)", 2)]),
    dict(id="US-17", epic="E9 Backend Terminal Core", layer="Backend", role="Backend",
         title="Terminal application core",
         acc="Role-based console menus; salted SHA-256 login; RBAC; console recovery; robust input handling.",
         prio="Must", pts=13, status="To Do", wf=[],
         subs=[("Console menu framework (role-based menus, clean exit)", 5),
               ("Console login with salted SHA-256 (parity with frontend)", 5),
               ("RBAC enforcement per menu action", 3),
               ("Console account recovery (3 questions, lockout)", 5),
               ("Input validation & error handling", 3)]),
    dict(id="US-18", epic="E10 Backend Domain & Audit", layer="Backend", role="Backend/Admin",
         title="Admin backend operations",
         acc="Console CRUD for clinics/staff with clinic-code generation and KPI aggregate queries.",
         prio="Must", pts=8, status="To Do", wf=[],
         subs=[("Enroll clinic & generate clinic code (insert)", 3),
               ("Onboard / deboard staff (insert/update + reassign)", 5),
               ("List clinics/staff & KPI aggregate queries", 3)]),
    dict(id="US-19", epic="E10 Backend Domain & Audit", layer="Backend", role="Backend/Doctor",
         title="Doctor backend operations",
         acc="Console read of schedule, upsert availability, insert document requests and consultation notes.",
         prio="Must", pts=8, status="To Do", wf=[],
         subs=[("View schedule & upsert availability", 5),
               ("Request document (insert, no vault access)", 3),
               ("Log consultation note (insert)", 2)]),
    dict(id="US-20", epic="E10 Backend Domain & Audit", layer="Backend", role="Backend/Receptionist",
         title="Receptionist backend operations",
         acc="List pending requests; assign a slot creating an appointment and updating the request atomically.",
         prio="Must", pts=5, status="To Do", wf=[],
         subs=[("List pending requests (query)", 2),
               ("Assign slot -> create appointment + update request (transaction)", 5)]),
    dict(id="US-21", epic="E10 Backend Domain & Audit", layer="Backend", role="Backend/Patient",
         title="Patient backend operations",
         acc="Submit appointment request (doctor guard), upload document metadata, grant/revoke time-boxed consent.",
         prio="Must", pts=8, status="To Do", wf=[],
         subs=[("Submit appointment request (guard: doctor exists)", 3),
               ("Upload document metadata (insert)", 2),
               ("Grant / revoke consent (time-boxed)", 5)]),
    dict(id="US-22", epic="E10 Backend Domain & Audit", layer="Backend", role="Backend",
         title="Audit logging & reporting",
         acc="Consent grant/revoke writes append-only audit rows; reporting queries print to console.",
         prio="Should", pts=5, status="To Do", wf=[],
         subs=[("Append-only audit_log on consent grant/revoke etc.", 3),
               ("Reporting queries printed to console (volume, roster)", 3)]),
]

wb = Workbook()

# =========================================================================== Cover
cov = wb.active
cov.title = "Cover"
cov["B2"] = "MediBridge"
cov["B2"].font = TITLE_FONT
cov["B3"] = "Clinical Operations Platform - Product Backlog"
cov["B3"].font = Font(bold=True, size=13, color=NAVY_MED)
cov["B5"] = "Architecture is delivered in two independent parts (not wired together):"
cov["B5"].font = BOLD
lines = [
    ("Frontend", "HTML + CSS + vanilla JavaScript. Standalone prototype using browser localStorage. Presented on its own."),
    ("Backend", "Vanilla Java + JDBC + PostgreSQL. TERMINAL / console application. Presented on its own."),
    ("Integration", "Out of scope - no API layer; frontend and backend/DB are demoed separately."),
    ("", ""),
    ("Sheets", "User Stories | Subtasks | Epics | Scope & Assumptions | Backend DB Schema"),
    ("Legend", "Priority = MoSCoW (Must/Should/Could).  Points = Fibonacci.  Status: Frontend = Done, Backend = To Do."),
]
r = 6
for a, b in lines:
    cov[f"B{r}"] = a
    cov[f"B{r}"].font = BOLD
    cov[f"C{r}"] = b
    cov[f"C{r}"].alignment = WRAP
    r += 1
tot = sum(s["pts"] for s in S)
n_sub = sum(len(s["subs"]) for s in S)
n_epics = len({s["epic"] for s in S})
cov[f"B{r+1}"] = "Totals"
cov[f"B{r+1}"].font = BOLD
cov[f"C{r+1}"] = f"{len(S)} user stories | {n_sub} subtasks | {n_epics} epics | {tot} story points"
cov.column_dimensions["A"].width = 3
cov.column_dimensions["B"].width = 16
cov.column_dimensions["C"].width = 95
cov.sheet_view.showGridLines = False

# =========================================================================== User Stories
ws = wb.create_sheet("User Stories")
headers = ["ID", "Epic", "Layer", "Role", "User Story", "Acceptance Criteria (key)",
           "Priority", "Points", "# Subtasks", "Status", "Wireframe(s)"]
ws.append(headers)
for s in S:
    ws.append([s["id"], s["epic"], s["layer"], s["role"], s["title"], s["acc"],
               s["prio"], s["pts"], len(s["subs"]), s["status"], ", ".join(s["wf"])])
style_header(ws, len(headers))
ws.freeze_panes = "A2"
widths = [8, 26, 11, 18, 34, 52, 9, 8, 11, 9, 40]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
for ridx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)), start=2):
    band = fill(BAND) if ridx % 2 == 0 else None
    for cell in row:
        cell.alignment = WRAP
        cell.border = BORDER
        if band:
            cell.fill = band
    row[2].alignment = CENTER
    row[2].fill = fill(LAYER_FILL.get(row[2].value, "FFFFFF"))
    row[6].alignment = CENTER
    row[6].fill = fill(PRIO_FILL.get(row[6].value, "FFFFFF"))
    row[7].alignment = CENTER
    row[8].alignment = CENTER
    row[9].alignment = CENTER
    row[9].fill = fill(STATUS_FILL.get(row[9].value, "FFFFFF"))
ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

# =========================================================================== Subtasks (grouped)
wst = wb.create_sheet("Subtasks")
sub_headers = ["ID", "User Story / Subtask", "Layer", "Priority", "Est (h)", "Status"]
wst.append(sub_headers)
wst.sheet_properties.outlinePr.summaryBelow = False
cur = 2
for s in S:
    wst.append([s["id"], s["title"], s["layer"], s["prio"], "", s["status"]])
    for cell in wst[cur]:
        cell.font = BOLD
        cell.fill = fill(EPIC_FILL_C)
        cell.alignment = WRAP
        cell.border = BORDER
    cur += 1
    for i, (text, est) in enumerate(s["subs"], start=1):
        wst.append([f"{s['id']}-T{i}", "    " + text, s["layer"], "", est, s["status"]])
        for cell in wst[cur]:
            cell.alignment = WRAP
            cell.border = BORDER
        wst[cur][4].alignment = CENTER
        wst[cur][5].alignment = CENTER
        wst[cur][5].fill = fill(STATUS_FILL.get(s["status"], "FFFFFF"))
        wst.row_dimensions[cur].outline_level = 1
        cur += 1
style_header(wst, len(sub_headers))
wst.freeze_panes = "A2"
for i, w in enumerate([12, 68, 11, 10, 8, 10], start=1):
    wst.column_dimensions[get_column_letter(i)].width = w

# =========================================================================== Epics
ws2 = wb.create_sheet("Epics")
ws2.append(["Epic", "Layer", "Stories", "# Stories", "Points"])
order = []
for s in S:
    if s["epic"] not in order:
        order.append(s["epic"])
epic_layer = {}
for s in S:
    epic_layer.setdefault(s["epic"], set()).add("Frontend" if s["layer"] == "Frontend" else "Backend/DB")
agg = {}
ids = {}
for s in S:
    c, p = agg.get(s["epic"], (0, 0))
    agg[s["epic"]] = (c + 1, p + s["pts"])
    ids.setdefault(s["epic"], []).append(s["id"])
for epic in order:
    c, p = agg[epic]
    layer = "Frontend" if epic_layer[epic] == {"Frontend"} else "Backend/DB"
    ws2.append([epic, layer, ", ".join(ids[epic]), c, p])
style_header(ws2, 5)
ws2.freeze_panes = "A2"
for i, w in enumerate([32, 12, 30, 10, 8], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for ridx, row in enumerate(ws2.iter_rows(min_row=2, max_row=ws2.max_row, max_col=5), start=2):
    band = fill(BAND) if ridx % 2 == 0 else None
    for cell in row:
        cell.alignment = WRAP
        cell.border = BORDER
        if band:
            cell.fill = band
    row[3].alignment = CENTER
    row[4].alignment = CENTER

# =========================================================================== Scope
ws3 = wb.create_sheet("Scope & Assumptions")
notes = [
    ("Architecture (delivered separately)", ""),
    ("Frontend", "HTML + external CSS + vanilla JavaScript only. Standalone prototype; data in browser localStorage. Presented on its own - no backend calls."),
    ("Backend", "Vanilla Java + JDBC + PostgreSQL. TERMINAL / console application. Presented on its own - NOT connected to the frontend."),
    ("Integration", "OUT OF SCOPE: no REST/API layer; DB used only by the Java console app; frontend & backend demoed separately."),
    ("", ""),
    ("In scope", ""),
    ("Roles", "Institute Admin, Doctor, Receptionist, Patient (RBAC per role)."),
    ("Security", "SHA-256 salted hashing for passwords & security answers; 3-question recovery with lockout; consent-based document access."),
    ("Data model", "Users, clinics (clinic codes), staff, patients, appointments & requests, availability, documents, consents/grants, audit log."),
    ("Delivery order", "Frontend built first (done); backend terminal app + Postgres built later."),
    ("", ""),
    ("Out of scope (explicitly)", ""),
    ("Clinical clutter", "No ICD-10 billing codes, medical licensing numbers, raw vitals tables, or medication dispensing."),
    ("Platform", "No mobile-native app; desktop-first web (1280-1440px). No live frontend-backend-DB connection."),
    ("Infra", "No cloud deployment, CI/CD, or real email/SMS delivery in this scope."),
    ("", ""),
    ("Estimation legend", ""),
    ("Priority", "MoSCoW - Must / Should / Could."),
    ("Points", "Fibonacci story points (relative effort)."),
    ("Subtask Est (h)", "Rough hours per subtask for task-level planning."),
    ("Status", "Frontend = Done (prototype built). Backend = To Do (later phase)."),
]
ws3.append(["Topic", "Detail"])
for tt, d in notes:
    ws3.append([tt, d])
style_header(ws3, 2)
ws3.freeze_panes = "A2"
ws3.column_dimensions["A"].width = 30
ws3.column_dimensions["B"].width = 95
for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, max_col=2):
    for cell in row:
        cell.alignment = WRAP
        cell.border = BORDER
    if row[1].value == "" and row[0].value:
        row[0].font = BOLD
        row[0].fill = fill(EPIC_FILL_C)
        row[1].fill = fill(EPIC_FILL_C)

# =========================================================================== DB schema
ws4 = wb.create_sheet("Backend DB Schema")
ws4.append(["Table", "Key Columns", "Notes"])
schema = [
    ("users", "id PK, full_name, email UNIQUE, password_hash, salt, role, clinic_id FK, status, created_at", "All roles incl. admin; role drives RBAC"),
    ("security_questions", "id PK, user_id FK, question, answer_hash", "3 rows per user; answers hashed"),
    ("clinics", "id PK, name, specialty, address, clinic_code UNIQUE, status, created_at", "clinic_code used at staff signup"),
    ("patients", "id PK, user_id FK, mrn UNIQUE, dob, blood_group, phone, allergies, emergency_contact", "Patient profile / digital pass data"),
    ("availability", "id PK, doctor_id FK, weekday, shift_start, shift_end, break_start, break_end", "Doctor weekly availability"),
    ("appointment_requests", "id PK, patient_id FK, doctor_id FK, preferred_date, reason, status, created_at", "status: pending/scheduled/declined"),
    ("appointments", "id PK, request_id FK, patient_id FK, doctor_id FK, clinic_id FK, date, time, visit_type, status", "Created when receptionist assigns a slot"),
    ("consultations", "id PK, appointment_id FK, doctor_id FK, patient_id FK, note, created_at", "Consultation history notes"),
    ("documents", "id PK, patient_id FK, file_name, mime_type, size_bytes, uploaded_at", "Metadata only in backend (no blob wiring)"),
    ("document_requests", "id PK, doctor_id FK, patient_id FK, message, status, created_at", "Privacy-preserving doc request"),
    ("consents", "id PK, patient_id FK, doctor_id FK, document_id FK, granted_at, expires_at, revoked_at", "Granular, time-boxed access grants"),
    ("audit_log", "id PK, actor_user_id, action, entity, entity_id, detail, created_at", "Append-only; consent grant/revoke etc."),
]
for x in schema:
    ws4.append(list(x))
style_header(ws4, 3)
ws4.freeze_panes = "A2"
for i, w in enumerate([22, 70, 40], start=1):
    ws4.column_dimensions[get_column_letter(i)].width = w
for ridx, row in enumerate(ws4.iter_rows(min_row=2, max_row=ws4.max_row, max_col=3), start=2):
    band = fill(BAND) if ridx % 2 == 0 else None
    for cell in row:
        cell.alignment = WRAP
        cell.border = BORDER
        if band:
            cell.fill = band
    row[0].font = BOLD

out = "docs/MediBridge_Product_Backlog.xlsx"
wb.save(out)
print(f"Saved {out} | {len(S)} stories, {n_sub} subtasks, {n_epics} epics, {tot} points")
